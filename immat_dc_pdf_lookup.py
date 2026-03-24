import os
import sys
import openpyxl


SCAN_ROOTS = {
    "AUTO":      r"\\srvfile03\FICHIERS-ERP\ATEMO-2021\scan",
    "INDUSTRIE": r"\\srvfile03\FICHIERS-ERP\ATEMO-INDUSTRIE\Scan",
    "MOTO":      r"\\srvfile06\DONNEES\Bases\Base_Surplus_Motos\Scan",
}


def build_pdf_index(root_path: str) -> dict[str, list[str]]:
    """Parcourt le répertoire une seule fois et indexe les PDFs par préfixe (premiers 7 caractères)."""
    index = {}
    total_files = 0
    print(f"Indexation des PDFs dans {root_path} ...", flush=True)
    for dirpath, dirnames, filenames in os.walk(root_path):
        pdfs = [f for f in filenames if f.lower().endswith(".pdf")]
        if pdfs:
            print(f"  -> {dirpath} ({len(pdfs)} PDF(s))", flush=True)
        for filename in pdfs:
            key = filename[:7].lower()
            full_path = os.path.join(dirpath, filename)
            index.setdefault(key, []).append(full_path)
            total_files += 1
    print(f"Index construit : {total_files} PDF(s) trouvés.\n", flush=True)
    return index


EXCLUDE_PATTERNS = [
    "_be", " be", "csa", "certcession", "cession", "perte", "achat", "fiv",
    "_re", "vies", "grais", "lettre", "offre", "rapport",
    "photos", "cmr", "affaires personnelles", "_pv", " pv", "mad", "frais",
    " da", "_da", "mise en demeure", " ci", "_ci", "bon", "positaire",
    "darva", "assurance", "facture", "vroom", "document", "attestation",
    "confirmation", "avis", "fiche", "enlevement", "lar", "cni",
]


def exclude_files(paths: list[str]) -> list[str]:
    result = []
    for p in paths:
        name_lower = os.path.basename(p).lower()
        # Pour _re : vérifier hors des occurrences de _recepisse
        name_for_re_check = name_lower.replace("_recepisse", "")
        excluded = False
        for pattern in EXCLUDE_PATTERNS:
            if pattern == "_re":
                if "_re" in name_for_re_check:
                    excluded = True
                    break
            elif pattern in name_lower:
                excluded = True
                break
        if not excluded:
            result.append(p)
    return result


def find_pdfs_in_index(index: dict, prefix: str) -> list[str]:
    prefix_lower = prefix.lower().strip()
    key = prefix_lower[:7]
    candidates = index.get(key, [])
    return [p for p in candidates if os.path.basename(p).lower().startswith(prefix_lower)]


def process_sheet(ws, index):
    ws.cell(row=1, column=10).value = "Nom du Fichier"
    ws.cell(row=1, column=11).value = "Cas"
    ws.cell(row=1, column=12).value = "Nb Résultats"
    updated = 0
    not_found = 0
    skipped = 0

    total = ws.max_row - 1
    print(f"Traitement de l'onglet {ws.title} ({total} lignes)...\n", flush=True)

    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        immat = row[3].value  # colonne D

        if not immat:
            skipped += 1
            continue

        immat_str = str(immat).strip()
        current = i - 1
        pct = current * 100 // total
        print(f"  [{pct:3d}%] ({current}/{total}) {immat_str} ...", end=" ", flush=True)

        all_results = exclude_files(find_pdfs_in_index(index, immat_str))
        prefix = immat_str.lower()
        dc_recepisse = [
            p for p in all_results
            if os.path.basename(p).lower().startswith(prefix + "_dc_recepisse")
        ]
        dc_only = [
            p for p in all_results
            if os.path.basename(p).lower().startswith(prefix + "_dc")
            or os.path.basename(p).lower().startswith(prefix + "_doc_decl_")
        ]
        if dc_recepisse:
            j_results = dc_recepisse
        elif dc_only:
            j_results = dc_only
        else:
            j_results = all_results  # fallback liste complète

        cell_j = ws.cell(row=i, column=10)  # colonne J : chemins complets
        cell_k = ws.cell(row=i, column=11)  # colonne K : noms de fichiers
        cell_l = ws.cell(row=i, column=12)  # colonne L : nb résultats

        if all_results:
            cell_j.value = " | ".join(j_results)
            cell_k.value = " | ".join(os.path.basename(p) for p in j_results)
            cell_l.value = len(j_results)
            updated += 1
            note = f"DC/DOC_DECL: {len(dc_only)}, total: {len(all_results)}"
            print(f"OK -> {note}", flush=True)
        else:
            cell_j.value = None
            cell_k.value = None
            cell_l.value = 0
            not_found += 1
            print("aucun fichier trouvé", flush=True)

    print(f"\nTerminé {ws.title}. {updated} ligne(s) mises à jour, {not_found} sans résultat, {skipped} ligne(s) vides ignorées.\n", flush=True)


def process_excel(excel_path: str):
    wb = openpyxl.load_workbook(excel_path)

    for sheet_name, scan_root in SCAN_ROOTS.items():
        if sheet_name not in wb.sheetnames:
            print(f"Onglet {sheet_name} introuvable, ignoré.\n", flush=True)
            continue

        if not os.path.exists(scan_root):
            print(f"Erreur: le chemin '{scan_root}' est inaccessible, onglet {sheet_name} ignoré.\n", flush=True)
            continue

        index = build_pdf_index(scan_root)
        process_sheet(wb[sheet_name], index)

    wb.save(excel_path)
    print("Fichier Excel sauvegardé.", flush=True)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python search_pdf_excel.py <chemin_vers_excel>")
        sys.exit(1)

    excel_path = sys.argv[1]

    if not os.path.exists(excel_path):
        print(f"Erreur: fichier Excel '{excel_path}' introuvable.")
        sys.exit(1)

    process_excel(excel_path)
